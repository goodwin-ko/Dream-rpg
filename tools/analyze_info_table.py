import openpyxl
import json

wb = openpyxl.load_workbook('d:/code/Dream/정보Table.xlsx', data_only=True)

with open('d:/code/Dream/data/game_data.json', 'r', encoding='utf-8') as f:
    current_game_data = json.load(f)

out = []

# --- 1. 보스드랍 Sheet Analysis ---
s_boss = wb['보스드랍']
boss_rows = list(s_boss.iter_rows(values_only=True))
boss_headers = boss_rows[0]
boss_data = boss_rows[1:]

boss_items = set()
boss_names = set()
item_types = set()
item_grades = set()

for r in boss_data:
    if not any(r): continue
    bname, iname, itype, igrade = r[0], r[1], r[2], r[3]
    if bname: boss_names.add(str(bname).strip())
    if iname: boss_items.add(str(iname).strip())
    if itype: item_types.add(str(itype).strip())
    if igrade: item_grades.add(str(igrade).strip())

out.append(f"=== [보스드랍 분석] ===")
out.append(f"총 행 수: {len(boss_data)}")
out.append(f"고유 보스 수: {len(boss_names)}")
out.append(f"고유 드랍 아이템 수: {len(boss_items)}")
out.append(f"아이템 종류 목록: {sorted(list(item_types))}")
out.append(f"장비 등급 목록: {sorted(list(item_grades))}")

current_bosses = {b['name'] for b in current_game_data.get('bosses', [])}
new_bosses = boss_names - current_bosses
out.append(f"현재 시스템에 없는 신규 보스 ({len(new_bosses)}개): {sorted(list(new_bosses))[:15]}...")

# --- 2. 장비아이템 Sheet Analysis ---
s_gear = wb['장비아이템']
gear_rows = list(s_gear.iter_rows(values_only=True))
gear_headers = gear_rows[0]
gear_data = gear_rows[1:]

gear_categories = set()
gear_grades = set()
items_with_effects = 0
items_with_active_passive = 0

sample_items = []

for r in gear_data:
    if not any(r): continue
    iname, cat, grade, req_lvl, opt_eff, act_pass = r[0], r[1], r[2], r[3], r[4], r[5]
    if cat: gear_categories.add(str(cat).strip())
    if grade: gear_grades.add(str(grade).strip())
    if opt_eff: items_with_effects += 1
    if act_pass: items_with_active_passive += 1
    if len(sample_items) < 5 and act_pass:
        sample_items.append({
            "name": iname, "cat": cat, "grade": grade, "lvl": req_lvl,
            "opt": opt_eff, "act_pass": act_pass
        })

out.append(f"\n=== [장비아이템 분석] ===")
out.append(f"총 장비 행 수: {len(gear_data)}")
out.append(f"장비 종류 목록: {sorted(list(gear_categories))}")
out.append(f"장비 등급 목록: {sorted(list(gear_grades))}")
out.append(f"옵션·전용효과 있는 장비 수: {items_with_effects}")
out.append(f"액티브·패시브 상세효과 있는 장비 수: {items_with_active_passive}")
out.append(f"\n[상세 효과 샘플 (액티브/패시브 포함)]:")
for s in sample_items:
    out.append(f"  * {s['name']} [{s['cat']} / {s['grade']} / Lv.{s['lvl']}]:")
    out.append(f"    - 기본옵션: {repr(s['opt'])}")
    out.append(f"    - 액티브/패시브: {repr(s['act_pass'])}")

# --- 3. 조합법 Sheet Analysis ---
s_recipe = wb['조합법']
rec_rows = list(s_recipe.iter_rows(values_only=True))
rec_headers = rec_rows[0]
rec_data = rec_rows[1:]

rec_items = set()
rec_cats = set()
rec_grades = set()

recipe_dict = {}
for r in rec_data:
    if not any(r): continue
    iname, cat, grade, mat_name, mat_qty = r[0], r[1], r[2], r[3], r[4]
    if iname:
        iname_s = str(iname).strip()
        rec_items.add(iname_s)
        if iname_s not in recipe_dict:
            recipe_dict[iname_s] = {"cat": cat, "grade": grade, "mats": []}
        if mat_name:
            recipe_dict[iname_s]["mats"].append({"name": str(mat_name).strip(), "qty": mat_qty})
    if cat: rec_cats.add(str(cat).strip())
    if grade: rec_grades.add(str(grade).strip())

out.append(f"\n=== [조합법 분석] ===")
out.append(f"총 조합법 행 수: {len(rec_data)}")
out.append(f"완성 아이템 수: {len(rec_items)}")
out.append(f"조합법 종류 목록: {sorted(list(rec_cats))}")
out.append(f"조합법 등급 목록: {sorted(list(rec_grades))}")

current_recipes = set(current_game_data.get('global_recipe_map', {}).keys())
new_recipes = rec_items - current_recipes
out.append(f"현재 시스템에 없는 신규 조합 아이템 ({len(new_recipes)}개): {sorted(list(new_recipes))[:15]}...")

with open('d:/code/Dream/tools/analyze_output.txt', 'w', encoding='utf-8') as f:
    f.write("\n".join(out))

print("Analysis complete! Output saved to tools/analyze_output.txt")
