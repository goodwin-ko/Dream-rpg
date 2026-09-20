import openpyxl
import os

excel_path = 'd:/code/Dream/정보Table.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)

out = []
out.append(f"Sheet names: {wb.sheetnames}\n")

for name in wb.sheetnames:
    sheet = wb[name]
    out.append(f"=== Sheet: {name} (Rows: {sheet.max_row}, Cols: {sheet.max_column}) ===")
    headers = [sheet.cell(row=1, column=c).value for c in range(1, sheet.max_column + 1)]
    out.append(f"Headers: {headers}")
    for r in range(2, min(sheet.max_row + 1, 8)):
        row_vals = [sheet.cell(row=r, column=c).value for c in range(1, min(sheet.max_column + 1, len(headers) + 1))]
        out.append(f"Row {r}: {row_vals}")
    out.append("")

with open('d:/code/Dream/tools/inspect_output.txt', 'w', encoding='utf-8') as f:
    f.write("\n".join(out))

print("DONE. Written to tools/inspect_output.txt")
