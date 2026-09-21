import openpyxl

wb_info = openpyxl.load_workbook('d:/code/Dream/정보Table.xlsx', data_only=True)
wb_dream = openpyxl.load_workbook('d:/code/Dream/dream01.xlsx', data_only=True)

out = []

def search_item(item_name):
    out.append(f"=== Searching for: {item_name} ===")
    for ws in wb_info.worksheets:
        for r in ws.iter_rows(values_only=True):
            r_str = " ".join(str(c) for c in r if c is not None)
            if item_name in r_str:
                out.append(f"  [정보Table - {ws.title}]: {[c for c in r if c is not None]}")
    for ws in wb_dream.worksheets:
        for r in ws.iter_rows(values_only=True):
            r_str = " ".join(str(c) for c in r if c is not None)
            if item_name in r_str:
                out.append(f"  [dream01 - {ws.title}]: {[c for c in r if c is not None]}")

for item in ["홍마룡의 칼날", "홍마룡", "진 · 블래스터", "진-블래스터", "진·블래스터", "진 블래스터", "블래스터"]:
    search_item(item)

with open('d:/code/Dream/tools/check_mats.txt', 'w', encoding='utf-8') as f:
    f.write("\n".join(out))

print("DONE")
