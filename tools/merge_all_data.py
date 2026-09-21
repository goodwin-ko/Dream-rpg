import openpyxl
import json
import os
import re

BASE_DIR = 'd:/code/Dream'
DREAM_EXCEL = os.path.join(BASE_DIR, 'dream01.xlsx')
INFO_EXCEL = os.path.join(BASE_DIR, '정보Table.xlsx')

ALIAS_MAP = {
    "천공의 성스로운 갑옷": "천공의 성스러운 갑옷",
    "여혼 정화의 성목": "영혼 정화의 성목",
    "빛의 그레 일린": "빛의 그레일린",
    "빛의 그러일런": "빛의 그레일린",
    "얼음 그림자 아퀄리스": "얼음 그림자 아퀼리스",
    "속세의 및-우키요에 두루마리": "속세의 빛·우키요에 두루마리",
    "속세의및-우키요에 두루마리": "속세의 빛·우키요에 두루마리",
    "속세의 및·우키요에 두루마리": "속세의 빛·우키요에 두루마리",
    "속세의및·우키요에 두루마리": "속세의 빛·우키요에 두루마리",
    "피의문양": "피의 문양",
    "영혼정화의 성목": "영혼 정화의 성목",
    "암흑의 회환": "암흑의 회한",
    "별의정령 알갈론": "별의정령 알갈론",
    "별의 정령 알갈론": "별의정령 알갈론",
    "빛의 정령 알갈론": "별의정령 알갈론",
    "별의 정령 알칼론": "별의정령 알갈론",
    "전생의 천사 임페리우스": "전쟁의 천사 임페리우스",
    "전쟁의 전사 임페리우스": "전쟁의 천사 임페리우스",
    "셩역의 대사제": "성역의 대사제",
    "여왕 아즈사라": "아즈샤라 여왕",
    "여왕 아즈샤라": "아즈샤라 여왕",
    "작열하는 용 이그닐": "작열하는 용 이그니르",
    "빙하의 죽음벌래 형제": "빙하의 죽음 벌레",
    "빙하의 죽음벌레 형제": "빙하의 죽음 벌레",
    "빙하의 죽음벌레": "빙하의 죽음 벌레",
    "마리안느": "저주받은 인형 마리안느",
    "실프": "광풍의 실프",
    "화염의 만샤": "광염의 만샤",
    "망령 메르겔": "망령 메르켈",
    "혈마의 시체": "피의 시체",
    "용망의 눈": "욕망의 눈",
    "명하의 수정 가지": "명해의 수정 가지",
    "점염마의 피": "적염마의 피",
    "열세의 화염깃": "멸세의 화염깃",
    "얼음 환상 바닞": "얼음 환상 반지",
    "결울 빙모": "겨울 빙모",
    "만샤의 영혼 햇불": "만샤의 영혼 횃불",
    "혼돈의 결정": "혼돈의 혈정",
    "성령의 조각": "성령석 조각",
    "지옥볼의 핵": "지옥불의 핵",
    "파천의뇌석": "파천의 뇌석",
    "빅옹의 결정": "빙옥의 결정",
    "끝없는 어룸의 로브": "끝없는 어둠의 로브",
    "프로토스 왕관": "프로스트 왕관",
    "지옥투구": "지옥 투구",
    "혼동의 혈정": "혼돈의 혈정",
    "영혼 녹이는 화염 결정": "혼을 녹이는 화염 결정",
    "성스러운 및의 수호석": "성스러운 빛의 수호석",
    "영혼 구술의 족쇄": "영혼 구슬의 족쇄",
    "용의 격노-용혼갑옷": "용의 격노·용혼 갑옷",
    "용의 격노-용혼 갑옷": "용의 격노·용혼 갑옷",
    "용의 격노·용혼갑옷": "용의 격노·용혼 갑옷",
    "잊힌 고대의 마나석": "잊힌 고대 마나석",
    "검은 쌍둥이 반지-성광 파편": "검은 쌍둥이 반지·섬광 파편",
    "검은 쌍둥이 반지·성광 파편": "검은 쌍둥이 반지·섬광 파편",
    "검은 쌍둥이 반지-섬광 파편": "검은 쌍둥이 반지·섬광 파편",
    "붉은 화염 전투 세트": "붉은 화염 전투 셋트",
    "격노한 화염의 전쟁신 도끼": "격노한 화염의 전생신 도끼",
    "청공의 분노 뇌전홀": "천공의 분노 뇌전홀",
    "TitanX-4060": "TitonX-4060",
    "성역의 대주교": "성역의 대사제",
    "용살자(근접)": "용살자 (근접)",
    "용살자 [근접]": "용살자 (근접)",
    "용살자[근접]": "용살자 (근접)",
    "용살자(원거리)": "용살자 (원거리)",
    "용살자 [원거리]": "용살자 (원거리)",
    "용살자[원거리]": "용살자 (원거리)"
}

def clean_str(s):
    if s is None:
        return ""
    val = str(s).strip()
    return ALIAS_MAP.get(val, val)

def normalize_name(s):
    if s is None:
        return ""
    val = str(s).strip()
    if val in ALIAS_MAP:
        val = ALIAS_MAP[val]
    if re.match(r'^tit[ao]nx-4060$', val, re.IGNORECASE):
        return "TitonX-4060"
    val = re.sub(r'\s*[\-\–—•ㆍ・]\s*', '·', val)
    val = re.sub(r'\s*·\s*', '·', val)
    if val in ALIAS_MAP:
        val = ALIAS_MAP[val]
    return val

def clean_lvl(s):
    if not s:
        return ""
    val = str(s).strip()
    m = re.search(r'\d+', val)
    return m.group(0) if m else val

def map_category(raw_cat):
    if not raw_cat:
        return "기타", ""
    c = str(raw_cat).strip()
    if c in ["갑옷", "경갑", "중갑", "천 갑옷", "천", "세트"]:
        sub = "중갑" if "중갑" in c else ("경갑" if "경갑" in c else ("천" if "천" in c else "갑옷"))
        return "갑옷", sub
    elif c in ["근접 무기", "근접무기"]:
        return "무기", "근접무기"
    elif c in ["원거리 무기", "원거리무기"]:
        return "무기", "원거리무기"
    elif c in ["마법 무기", "마법무기", "지팡이"]:
        return "무기", "지팡이"
    elif c in ["무기"]:
        return "무기", "무기"
    elif c in ["투구", "머리장식"]:
        sub = "머리장식" if "머리장식" in c else "투구"
        return "투구", sub
    elif c in ["장신구"]:
        return "장신구", "장신구"
    elif c in ["보조 장비", "보조장비", "배지"]:
        sub = "배지" if "배지" in c else "보조장비"
        return "보조장비", sub
    elif c in ["마나석"]:
        return "마나석", "마나석"
    return c, c

print("Importing parse_excel from tools...")
import sys
sys.path.insert(0, os.path.join(BASE_DIR, 'tools'))
from parse_excel import parse_excel

print("1. Parsing dream01.xlsx base...")
data = parse_excel(DREAM_EXCEL)
boss_map = {b["name"]: b for b in data["bosses"]}
global_recipe_map = data["global_recipe_map"]
category_gear = data["category_gear"]

# Load 정보Table.xlsx
print("2. Loading 정보Table.xlsx...")
wb_info = openpyxl.load_workbook(INFO_EXCEL, data_only=True)

# 2-1. Process 장비아이템 (Gear Specs)
print("3. Processing 장비아이템...")
s_gear = wb_info["장비아이템"]
gear_specs = {}

for r in list(s_gear.iter_rows(values_only=True))[1:]:
    if not any(r): continue
    iname = normalize_name(r[0])
    raw_cat = clean_str(r[1])
    grade = clean_str(r[2])
    lvl = r[3] if r[3] is not None else 0
    opt = str(r[4]).strip() if r[4] else ""
    act_pass = str(r[5]).strip() if r[5] else ""
    
    cat, sub_type = map_category(raw_cat)
    lvl_num = int(lvl) if str(lvl).isdigit() else 0
    
    if iname == "용살자":
        if sub_type == "근접무기" or lvl_num == 360:
            iname = "용살자 (근접)"
        elif sub_type == "원거리무기" or lvl_num == 320:
            iname = "용살자 (원거리)"

    gear_specs[iname] = {
        "name": iname,
        "category": cat,
        "sub_type": sub_type,
        "grade": grade,
        "level": lvl_num,
        "level_str": f"Lv.{lvl_num}" if lvl_num else "",
        "options": opt,
        "active_passive": act_pass
    }

print(f"Loaded {len(gear_specs)} gear specs from 장비아이템")

# 2-2. Process 조합법
print("4. Processing 조합법...")
s_rec = wb_info["조합법"]
rec_groups = {}
for r in list(s_rec.iter_rows(values_only=True))[1:]:
    if not any(r): continue
    iname = normalize_name(r[0])
    raw_cat = clean_str(r[1])
    grade = clean_str(r[2])
    mat_name = normalize_name(r[3])
    mat_qty = r[4] if r[4] is not None else 1
    
    if not iname: continue
    cat, sub_type = map_category(raw_cat)
    
    if iname == "용살자":
        if sub_type == "근접무기":
            iname = "용살자 (근접)"
        elif sub_type == "원거리무기":
            iname = "용살자 (원거리)"

    if mat_name == "용살자":
        if iname == "황룡언월도" or sub_type == "근접무기":
            mat_name = "용살자 (근접)"
        elif iname == "파괴자" or sub_type == "원거리무기":
            mat_name = "용살자 (원거리)"

    if iname not in rec_groups:
        rec_groups[iname] = {
            "name": iname,
            "category": cat,
            "sub_type": sub_type,
            "grade": grade,
            "materials": []
        }
    if mat_name:
        try:
            qty_num = int(mat_qty)
        except:
            qty_num = 1
        rec_groups[iname]["materials"].append({
            "name": mat_name,
            "qty": qty_num
        })

print(f"Loaded {len(rec_groups)} recipes from 조합법")

# Merge 조합법 into global_recipe_map and category_gear
for iname, rinfo in rec_groups.items():
    gspec = gear_specs.get(iname, {})
    cat = gspec.get("category") or rinfo["category"]
    sub_type = gspec.get("sub_type") or rinfo["sub_type"]
    grade = gspec.get("grade") or rinfo["grade"]
    lvl = gspec.get("level") or 0
    lvl_str = gspec.get("level_str") or (f"Lv.{lvl}" if lvl else "")
    opt = gspec.get("options") or ""
    act_pass = gspec.get("active_passive") or ""

    if iname in global_recipe_map:
        rc = global_recipe_map[iname]
        rc["grade"] = grade
        if opt: rc["options"] = opt
        if act_pass: rc["active_passive"] = act_pass
        if lvl and not rc.get("level"):
            rc["level"] = lvl
            rc["level_str"] = lvl_str
        if sub_type and not rc.get("sub_cat"):
            rc["sub_cat"] = sub_type
        # If existing materials was empty, or for key items to ensure complete materials, use new materials
        if (not rc.get("materials") or iname in ["용살자 (근접)", "용살자 (원거리)", "황룡언월도", "파괴자"]) and rinfo["materials"]:
            rc["materials"] = rinfo["materials"]
    else:
        # New recipe!
        global_recipe_map[iname] = {
            "name": iname,
            "category": cat,
            "sub_cat": sub_type,
            "grade": grade,
            "level": lvl,
            "level_str": lvl_str,
            "synergy": "",
            "synergy_jobs": [],
            "special_effect": "",
            "options": opt,
            "active_passive": act_pass,
            "materials": rinfo["materials"],
            "is_drop": False,
            "drop_boss": "",
            "drop_level": 0,
            "drop_location": "",
            "drop_type": "아이템"
        }
        if cat in category_gear:
            if iname not in category_gear[cat]["top_gear"]:
                category_gear[cat]["top_gear"].append(iname)
            category_gear[cat]["recipes"].append(global_recipe_map[iname])

# Enrich any other recipes with gear_specs
for iname, gspec in gear_specs.items():
    if iname in global_recipe_map:
        rc = global_recipe_map[iname]
        if gspec.get("grade") and not rc.get("grade"):
            rc["grade"] = gspec["grade"]
        if gspec.get("options") and not rc.get("options"):
            rc["options"] = gspec["options"]
        if gspec.get("active_passive") and not rc.get("active_passive"):
            rc["active_passive"] = gspec["active_passive"]
        if gspec.get("level") and not rc.get("level"):
            rc["level"] = gspec["level"]
            rc["level_str"] = gspec["level_str"]
        if gspec.get("sub_type") and not rc.get("sub_cat"):
            rc["sub_cat"] = gspec["sub_type"]
    else:
        # Item has gear specs but not craftable (e.g. drop item or uncrafted top gear)
        cat = gspec["category"]
        if cat in category_gear:
            global_recipe_map[iname] = {
                "name": iname,
                "category": cat,
                "sub_cat": gspec["sub_type"],
                "grade": gspec["grade"],
                "level": gspec["level"],
                "level_str": gspec["level_str"],
                "synergy": "",
                "synergy_jobs": [],
                "special_effect": "",
                "options": gspec["options"],
                "active_passive": gspec["active_passive"],
                "materials": [],
                "is_drop": True,
                "drop_boss": "",
                "drop_level": gspec["level"],
                "drop_location": "",
                "drop_type": "아이템"
            }
            if iname not in category_gear[cat]["top_gear"]:
                category_gear[cat]["top_gear"].append(iname)
            category_gear[cat]["recipes"].append(global_recipe_map[iname])

# 2-3. Process 보스드랍
print("5. Processing 보스드랍...")
s_boss = wb_info["보스드랍"]
for r in list(s_boss.iter_rows(values_only=True))[1:]:
    if not any(r): continue
    bname = clean_str(r[0])
    iname = normalize_name(r[1])
    itype = clean_str(r[2])
    igrade = clean_str(r[3])
    
    if not bname or not iname: continue
    
    if iname == "용살자":
        iname = "용살자 (원거리)"
    
    # If boss not in boss_map, add
    if bname not in boss_map:
        boss_map[bname] = {
            "name": bname,
            "level": 0,
            "level_str": "",
            "location": "",
            "drops": []
        }
    
    # Check if iname is already dropped by bname
    existing_drops = boss_map[bname]["drops"]
    drop_entry = next((d for d in existing_drops if d["name"] == iname), None)
    if drop_entry:
        if itype and not drop_entry.get("type"):
            drop_entry["type"] = itype
        if igrade and not drop_entry.get("grade"):
            drop_entry["grade"] = igrade
    else:
        # Add new drop
        new_drop = {
            "name": iname,
            "type": itype or "아이템",
            "grade": igrade or (gear_specs.get(iname, {}).get("grade", "")),
            "level": boss_map[bname]["level"],
            "level_str": boss_map[bname]["level_str"]
        }
        boss_map[bname]["drops"].append(new_drop)
        
    # Update recipe drop info if item in global_recipe_map
    if iname in global_recipe_map:
        rc = global_recipe_map[iname]
        if not rc.get("drop_boss"):
            rc["drop_boss"] = bname
            rc["drop_location"] = boss_map[bname].get("location", "")
        if igrade and not rc.get("grade"):
            rc["grade"] = igrade

# Ensure 용살자 disambiguation in global_recipe_map and category_gear
if "용살자 (근접)" in global_recipe_map:
    r_melee = global_recipe_map["용살자 (근접)"]
    r_melee["category"] = "무기"
    r_melee["sub_cat"] = "근접무기"
    r_melee["grade"] = r_melee.get("grade") or "전설"
    r_melee["level"] = 360
    r_melee["level_str"] = "Lv.360"
    r_melee["synergy"] = "광전사"
    r_melee["synergy_jobs"] = ["광전사"]
    r_melee["special_effect"] = "광전사"
    if "용살자 (근접)" in rec_groups and rec_groups["용살자 (근접)"]["materials"]:
        r_melee["materials"] = rec_groups["용살자 (근접)"]["materials"]

if "용살자 (원거리)" in global_recipe_map:
    r_ranged = global_recipe_map["용살자 (원거리)"]
    r_ranged["category"] = "무기"
    r_ranged["sub_cat"] = "원거리무기"
    r_ranged["grade"] = r_ranged.get("grade") or "에픽"
    r_ranged["level"] = 320
    r_ranged["level_str"] = "Lv.320"
    if "용살자 (원거리)" in rec_groups and rec_groups["용살자 (원거리)"]["materials"]:
        r_ranged["materials"] = rec_groups["용살자 (원거리)"]["materials"]

if "황룡언월도" in global_recipe_map:
    for m in global_recipe_map["황룡언월도"].get("materials", []):
        if m["name"] in ["용살자", "용살자 (원거리)"]:
            m["name"] = "용살자 (근접)"
            m["level"] = 360
            m["level_str"] = "Lv.360"

if "파괴자" in global_recipe_map:
    for m in global_recipe_map["파괴자"].get("materials", []):
        if m["name"] in ["용살자", "용살자 (근접)"]:
            m["name"] = "용살자 (원거리)"
            m["level"] = 320
            m["level_str"] = "Lv.320"

global_recipe_map.pop("용살자", None)

for cat, cg in category_gear.items():
    if "top_gear" in cg:
        new_top = []
        seen = set()
        for x in cg["top_gear"]:
            x_clean = ALIAS_MAP.get(x, x)
            if x_clean == "용살자":
                if "용살자 (근접)" not in seen:
                    seen.add("용살자 (근접)")
                    new_top.append("용살자 (근접)")
                if "용살자 (원거리)" not in seen:
                    seen.add("용살자 (원거리)")
                    new_top.append("용살자 (원거리)")
            else:
                if x_clean not in seen:
                    seen.add(x_clean)
                    new_top.append(x_clean)
        if cat == "무기":
            if "용살자 (근접)" not in seen:
                new_top.append("용살자 (근접)")
            if "용살자 (원거리)" not in seen:
                new_top.append("용살자 (원거리)")
        cg["top_gear"] = new_top

    if "recipes" in cg:
        rec_map = {}
        for rc in cg["recipes"]:
            rname = ALIAS_MAP.get(rc["name"], rc["name"])
            if rname == "용살자":
                continue
            rec_map[rname] = rc
        if cat == "무기":
            if "용살자 (근접)" in global_recipe_map:
                rec_map["용살자 (근접)"] = global_recipe_map["용살자 (근접)"]
            if "용살자 (원거리)" in global_recipe_map:
                rec_map["용살자 (원거리)"] = global_recipe_map["용살자 (원거리)"]
        cg["recipes"] = list(rec_map.values())

# 2-4. Resolve material attributes (boss, level, location, is_drop)
drop_lookup = {}
for b in boss_map.values():
    for d in b.get("drops", []):
        drop_lookup[d["name"]] = {
            "boss": b["name"],
            "level": d.get("level") or b.get("level") or 0,
            "level_str": d.get("level_str") or b.get("level_str") or "",
            "location": b.get("location", ""),
            "type": d.get("type", "아이템")
        }

for rc in global_recipe_map.values():
    for m in rc.get("materials", []):
        mname = m.get("name")
        if not mname: continue
        if not m.get("boss") or m.get("boss") == "조합템" or not m.get("level"):
            if mname in drop_lookup:
                dinfo = drop_lookup[mname]
                m["boss"] = dinfo["boss"]
                m["level"] = dinfo["level"]
                m["level_str"] = dinfo["level_str"]
                m["location"] = dinfo["location"]
                m["is_drop"] = True
            elif mname in global_recipe_map:
                sub_r = global_recipe_map[mname]
                if sub_r.get("is_drop") and sub_r.get("drop_boss"):
                    m["boss"] = sub_r["drop_boss"]
                    m["level"] = sub_r.get("drop_level") or sub_r.get("level") or 0
                    m["level_str"] = sub_r.get("level_str") or (f"Lv.{m['level']}" if m["level"] else "")
                    m["location"] = sub_r.get("drop_location", "")
                    m["is_drop"] = True
                else:
                    m["boss"] = "조합템"
                    m["level"] = sub_r.get("level") or 0
                    m["level_str"] = sub_r.get("level_str") or (f"Lv.{m['level']}" if m["level"] else "")

# 2-5. Rebuild 31 jobs using merged data
print("6. Rebuilding jobs with merged recipes...")
categories = ["무기", "갑옷", "투구", "장신구", "보조장비", "마나석"]
ALL_31_JOBS = data["all_job_list"]
jobs = {}
legacy_jobs = data.get("jobs", {})

for job_name in ALL_31_JOBS:
    recomms = {cat: [] for cat in categories}
    for cat in categories:
        for rc in category_gear[cat]["recipes"]:
            if job_name in rc.get("synergy_jobs", []):
                recomms[cat].append({
                    "name": rc["name"],
                    "level": rc["level"],
                    "synergy": rc["synergy"],
                    "sub_cat": rc["sub_cat"]
                })
                
    default_loadout = {}
    if job_name in legacy_jobs:
        for cat in categories:
            top_items = legacy_jobs[job_name]["gear_slots"].get(cat, [])
            if isinstance(top_items, list) and top_items:
                default_loadout[cat] = top_items[0]
            elif isinstance(top_items, str) and top_items:
                default_loadout[cat] = top_items
            elif category_gear[cat]["top_gear"]:
                default_loadout[cat] = category_gear[cat]["top_gear"][0]
            else:
                default_loadout[cat] = ""
    else:
        for cat in categories:
            cat_top = category_gear[cat]["top_gear"]
            job_recs = [r["name"] for r in recomms[cat] if r["name"] in cat_top]
            if job_recs:
                default_loadout[cat] = job_recs[0]
            elif cat_top:
                default_loadout[cat] = cat_top[0]
            else:
                default_loadout[cat] = ""
                
    gear_slots = {cat: [default_loadout[cat]] if default_loadout.get(cat) else [] for cat in categories}
    
    active_recipes = []
    visited = set()
    def collect_recipes(item_name):
        if not item_name or item_name in visited:
            return
        visited.add(item_name)
        rc = global_recipe_map.get(item_name)
        if rc:
            active_recipes.append(rc)
            for m in rc.get("materials", []):
                m_name = m.get("name")
                if m_name in global_recipe_map:
                    collect_recipes(m_name)
                    
    for slot_item in default_loadout.values():
        if slot_item:
            collect_recipes(slot_item)
            
    active_recipe_map = {r["name"]: r for r in active_recipes}
    
    jobs[job_name] = {
        "name": job_name,
        "recipes": active_recipes,
        "recipe_map": active_recipe_map,
        "gear_slots": gear_slots,
        "default_loadout": default_loadout,
        "recommendations": recomms
    }
data["jobs"] = jobs

# Re-sort bosses by level, name
sorted_bosses = sorted(list(boss_map.values()), key=lambda x: (x["level"], x["name"]))
data["bosses"] = sorted_bosses

# Save merged output to data/game_data.json and data/game_data.js
output_json = os.path.join(BASE_DIR, 'data', 'game_data.json')
output_js = os.path.join(BASE_DIR, 'data', 'game_data.js')

with open(output_json, 'w', encoding='utf-8') as f:
    json.dump(data, f, ensure_ascii=False, indent=2)

with open(output_js, 'w', encoding='utf-8') as f:
    f.write("window.PRELOADED_GAME_DATA = " + json.dumps(data, ensure_ascii=False, indent=2) + ";\n")

print(f"SUCCESS! Merged game data written to:")
print(f"  {output_json}")
print(f"  {output_js}")
print(f"Total bosses: {len(data['bosses'])}")
print(f"Total global recipes: {len(data['global_recipe_map'])}")
for cat in ["무기", "갑옷", "투구", "장신구", "보조장비", "마나석"]:
    print(f"  [{cat}]: {len(data['category_gear'][cat]['top_gear'])} top gear, {len(data['category_gear'][cat]['recipes'])} total recipes")
