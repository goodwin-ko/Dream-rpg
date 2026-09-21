import openpyxl

wb_info = openpyxl.load_workbook('d:/code/Dream/정보Table.xlsx', data_only=True)
wb_dream = openpyxl.load_workbook('d:/code/Dream/dream01.xlsx', data_only=True)

out = []

out.append("=== 파괴자 in dream01.xlsx ===")
ws = wb_dream['무기']
for r_idx, r in enumerate(ws.iter_rows(values_only=True), 1):
    r_str = " ".join(str(c) for c in r if c is not None)
    if "파괴자" in r_str:
        out.append(f"  Row {r_idx}: {[c for c in r if c is not None]}")

out.append("=== 파괴자 in 정보Table.xlsx ===")
ws2 = wb_info['조합법']
for r_idx, r in enumerate(ws2.iter_rows(values_only=True), 1):
    r_str = " ".join(str(c) for c in r if c is not None)
    if "파괴자" in r_str:
        out.append(f"  Row {r_idx}: {[c for c in r if c is not None]}")

with open('d:/code/Dream/tools/destroyer_info.txt', 'w', encoding='utf-8') as f:
    f.write("\n".join(out))

print("DONE")
