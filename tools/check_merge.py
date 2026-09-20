import openpyxl
import json

wb = openpyxl.load_workbook('d:/code/Dream/정보Table.xlsx', data_only=True)

with open('d:/code/Dream/data/game_data.json', 'r', encoding='utf-8') as f:
    current_data = json.load(f)

# Inspect recipes
s_rec = wb['조합법']
rec_rows = list(s_rec.iter_rows(values_only=True))[1:]
rec_dict = {}
for r in rec_rows:
    if not any(r): continue
    iname = str(r[0]).strip() if r[0] else ""
    if not iname: continue
    cat = str(r[1]).strip() if r[1] else ""
    grade = str(r[2]).strip() if r[2] else ""
    mat = str(r[3]).strip() if r[3] else ""
    qty = r[4] if r[4] is not None else 1
    if iname not in rec_dict:
        rec_dict[iname] = {"cat": cat, "grade": grade, "materials": []}
    if mat:
        rec_dict[iname]["materials"].append({"name": mat, "qty": qty})

# Inspect gear stats
s_gear = wb['장비아이템']
gear_rows = list(s_gear.iter_rows(values_only=True))[1:]
gear_stats = {}
for r in gear_rows:
    if not any(r): continue
    iname = str(r[0]).strip() if r[0] else ""
    if not iname: continue
    cat = str(r[1]).strip() if r[1] else ""
    grade = str(r[2]).strip() if r[2] else ""
    lvl = r[3] if r[3] is not None else 0
    opt = str(r[4]).strip() if r[4] else ""
    act_pass = str(r[5]).strip() if r[5] else ""
    gear_stats[iname] = {
        "cat": cat, "grade": grade, "level": lvl,
        "options": opt, "active_passive": act_pass
    }

# Inspect boss drops
s_boss = wb['보스드랍']
boss_rows = list(s_boss.iter_rows(values_only=True))[1:]
boss_drops_map = {}
for r in boss_rows:
    if not any(r): continue
    bname = str(r[0]).strip() if r[0] else ""
    iname = str(r[1]).strip() if r[1] else ""
    itype = str(r[2]).strip() if r[2] else ""
    igrade = str(r[3]).strip() if r[3] else ""
    if not bname or not iname: continue
    if bname not in boss_drops_map:
        boss_drops_map[bname] = []
    boss_drops_map[bname].append({"name": iname, "type": itype, "grade": igrade})

out = []
out.append(f"총 조합법 아이템 수: {len(rec_dict)}")
out.append(f"총 장비 스펙 아이템 수: {len(gear_stats)}")
out.append(f"총 보스 수: {len(boss_drops_map)}")

# Check how many recipes have gear stats
matched_gear = sum(1 for k in rec_dict if k in gear_stats)
out.append(f"조합법 아이템 중 장비 상세스펙(옵션/스킬) 매칭: {matched_gear}/{len(rec_dict)}")

# Check category distribution of gear items
cat_dist = {}
for k, v in gear_stats.items():
    cat = v['cat']
    cat_dist[cat] = cat_dist.get(cat, 0) + 1
out.append(f"\n장비아이템 시트 카테고리 분포: {cat_dist}")

with open('d:/code/Dream/tools/merge_stats.txt', 'w', encoding='utf-8') as f:
    f.write("\n".join(out))

print("Merge stats written to tools/merge_stats.txt")
