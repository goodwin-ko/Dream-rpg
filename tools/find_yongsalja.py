import openpyxl
import json
import os

out = []

def check_file(filepath, label):
    out.append(f"=== {label} ({filepath}) ===")
    if not os.path.exists(filepath):
        out.append("File not found")
        return
    wb = openpyxl.load_workbook(filepath, data_only=True)
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        matches = []
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            row_str = " ".join(str(c) for c in row if c is not None)
            if "용살자" in row_str or "황룡언월도" in row_str:
                matches.append(f"  Row {r_idx}: {[c for c in row if c is not None]}")
        if matches:
            out.append(f"[{sheetname}] ({len(matches)} matches):")
            out.extend(matches)

check_file('d:/code/Dream/dream01.xlsx', 'dream01.xlsx')
check_file('d:/code/Dream/정보Table.xlsx', '정보Table.xlsx')

with open('d:/code/Dream/tools/yongsalja_info.txt', 'w', encoding='utf-8') as f:
    f.write("\n".join(out))

print("DONE. Written to tools/yongsalja_info.txt")
