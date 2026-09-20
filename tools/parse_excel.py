import openpyxl
import json
import os
import re

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EXCEL_PATH = os.path.join(BASE_DIR, 'dream01.xlsx')
OUTPUT_JSON = os.path.join(BASE_DIR, 'data', 'game_data.json')
OUTPUT_JS = os.path.join(BASE_DIR, 'data', 'game_data.js')

# Normalization dictionary for typos / variations across all sheets
ALIAS_MAP = {
    "천공의 성스로운 갑옷": "천공의 성스러운 갑옷",
    "여혼 정화의 성목": "영혼 정화의 성목",
    "빛의 그레 일린": "빛의 그레일린",
    "빛의 그러일런": "빛의 그레일린",
    "얼음 그림자 아퀄리스": "얼음 그림자 아퀼리스",
    "속세의 및-우키요에 두루마리": "속세의 빛-우키요에 두루마리",
    "속세의및-우키요에 두루마리": "속세의 빛-우키요에 두루마리",
    "피의문양": "피의 문양",
    "영혼정화의 성목": "영혼 정화의 성목",
    "암흑의 회환": "암흑의 회한",
    "별의정령 알갈론": "별의정령 알갈론",
    "별의 정령 알갈론": "별의정령 알갈론",
    "빛의 정령 알갈론": "별의정령 알갈론",
    "별의 정령 알칼론": "별의정령 알갈론",
    "전생의 천사 임페리우스": "전쟁의 천사 임페리우스",
    "전쟁의 전사 임페리우스": "전쟁의 천사 임페리우스",
    "셩역의 대사제": "성역의 대사제",
    "여왕 아즈사라": "아즈샤라 여왕",
    "여왕 아즈샤라": "아즈샤라 여왕",
    "작열하는 용 이그닐": "작열하는 용 이그니르",
    "빙하의 죽음벌래 형제": "빙하의 죽음 벌레",
    "빙하의 죽음벌레 형제": "빙하의 죽음 벌레",
    "빙하의 죽음벌레": "빙하의 죽음 벌레",
    "마리안느": "저주받은 인형 마리안느",
    "실프": "광풍의 실프",
    "화염의 만샤": "광염의 만샤",
    "망령 메르겔": "망령 메르켈",
    "혈마의 시체": "피의 시체",
    "용망의 눈": "욕망의 눈",
    "명하의 수정 가지": "명해의 수정 가지",
    "점염마의 피": "적염마의 피",
    "열세의 화염깃": "멸세의 화염깃",
    "얼음 환상 바닞": "얼음 환상 반지",
    "결울 빙모": "겨울 빙모",
    "만샤의 영혼 햇불": "만샤의 영혼 횃불",
    "혼돈의 결정": "혼돈의 혈정",
    "성령의 조각": "성령석 조각",
    "지옥볼의 핵": "지옥불의 핵",
    "파천의뇌석": "파천의 뇌석",
    "빅옹의 결정": "빙옥의 결정",
    "끝없는 어룸의 로브": "끝없는 어둠의 로브",
    # Aliases for category sheets
    "프로토스 왕관": "프로스트 왕관",
    "지옥투구": "지옥 투구",
    "혼동의 혈정": "혼돈의 혈정",
    "영혼 녹이는 화염 결정": "혼을 녹이는 화염 결정",
    "성스러운 및의 수호석": "성스러운 빛의 수호석",
    "영혼 구술의 족쇄": "영혼 구슬의 족쇄",
    "용의 격노-용혼갑옷": "용의 격노-용혼 갑옷",
    "잊힌 고대의 마나석": "잊힌 고대 마나석",
    "검은 쌍둥이 반지-성광 파편": "검은 쌍둥이 반지-섬광 파편",
    "붉은 화염 전투 세트": "붉은 화염 전투 셋트",
    "격노한 화염의 전쟁신 도끼": "격노한 화염의 전생신 도끼",
    "청공의 분노 뇌전홀": "천공의 분노 뇌전홀",
}

def clean_str(s):
    if s is None:
        return ""
    val = str(s).strip()
    return ALIAS_MAP.get(val, val)

def clean_lvl(s):
    if not s:
        return ""
    val = str(s).strip()
    m = re.search(r'\d+', val)
    return m.group(0) if m else val

ALL_31_JOBS = [
    "네크로맨서", "인형술사", "점성술사", "프리스트", "얼음마법사",
    "어둠마법사", "비전마법사", "화염마법사", "드루이드", "번개마법사",
    "원소마법사", "암살자", "궁수", "검성", "사수",
    "블레이드 스피릿", "거너", "머스킷티어", "싸움꾼", "기계공",
    "권법가", "격투가", "가디언", "살육자", "검혼",
    "검사", "광전사", "마검사", "죽음의기사", "배틀메이지", "성기사"
]

def parse_excel(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    
    # 1. Parse Boss sheet with '이동지역' (Location)
    boss_map = {}
    item_to_boss = {}
    
    if "보스" in wb.sheetnames:
        s_boss = wb["보스"]
        rows = list(s_boss.iter_rows(values_only=True))
        for r in rows[1:]:
            if not any(r):
                continue
            b_name = clean_str(r[0])
            lvl = clean_lvl(r[1])
            it_name = clean_str(r[2])
            cat = clean_str(r[3]) if len(r) > 3 and r[3] is not None else "아이템"
            loc = clean_str(r[4]) if len(r) > 4 and r[4] is not None else ""
            
            if not b_name:
                continue
            lvl_num = int(lvl) if lvl.isdigit() else 0
            if b_name not in boss_map:
                boss_map[b_name] = {
                    "name": b_name,
                    "level": lvl_num,
                    "level_str": f"Lv.{lvl_num}" if lvl_num else "",
                    "location": loc,
                    "drops": []
                }
            if loc and not boss_map[b_name]["location"]:
                boss_map[b_name]["location"] = loc
                
            if it_name:
                boss_map[b_name]["drops"].append({
                    "name": it_name,
                    "type": cat,
                    "level": lvl_num,
                    "level_str": f"Lv.{lvl_num}" if lvl_num else ""
                })
                item_to_boss[it_name] = {
                    "boss": b_name,
                    "level": lvl_num,
                    "level_str": f"Lv.{lvl_num}" if lvl_num else "",
                    "location": loc,
                    "type": cat
                }

    # Propagate boss location to all items in item_to_boss
    for it_name, info in item_to_boss.items():
        b_name = info["boss"]
        if b_name in boss_map and boss_map[b_name]["location"] and not info["location"]:
            info["location"] = boss_map[b_name]["location"]

    # Fallback mappings for unlisted drops
    item_to_boss["이그닐의 심장"] = {"boss": "작열하는 용 이그니르", "level": 340, "level_str": "Lv.340", "location": boss_map.get("작열하는 용 이그니르", {}).get("location", "용암 화산-왼쪽포탈-오른쪽"), "type": "재료"}
    item_to_boss["이그니르의 심장"] = {"boss": "작열하는 용 이그니르", "level": 340, "level_str": "Lv.340", "location": boss_map.get("작열하는 용 이그니르", {}).get("location", "용암 화산-왼쪽포탈-오른쪽"), "type": "재료"}
    if "작열하는 용 이그니르" in boss_map:
        existing_drops = [d["name"] for d in boss_map["작열하는 용 이그니르"]["drops"]]
        if "이그니르의 심장" not in existing_drops and "이그닐의 심장" not in existing_drops:
            boss_map["작열하는 용 이그니르"]["drops"].append({"name": "이그니르의 심장", "type": "재료", "level": 340, "level_str": "Lv.340"})
            
    # Consolidated Mining Entry in Boss Drops
    mining_drops = [
        {"name": "철광석", "type": "채광", "location": "광산", "display": "철광석 - 광산", "level": 0, "level_str": "채광"},
        {"name": "은광석", "type": "채광", "location": "광산", "display": "은광석 - 광산", "level": 0, "level_str": "채광"},
        {"name": "금광석", "type": "채광", "location": "광산", "display": "금광석 - 광산", "level": 0, "level_str": "채광"},
        {"name": "맑은 샘의 보석", "type": "채광", "location": "멀록해변", "display": "맑은 샘의 보석 - 멀록해변", "level": 0, "level_str": "채광"},
        {"name": "현빙", "type": "채광", "location": "운설산", "display": "현빙 - 운설산", "level": 0, "level_str": "채광"},
        {"name": "지옥석", "type": "채광", "location": "용암화산", "display": "지옥석 - 용암화산", "level": 0, "level_str": "채광"},
    ]
    boss_map["채광"] = {
        "name": "채광",
        "level": 0,
        "level_str": "채광",
        "location": "광산 / 멀록해변 / 운설산 / 용암화산",
        "drops": mining_drops
    }
    
    # Map mining items in item_to_boss with specific zones
    item_to_boss["철광석"] = {"boss": "광산-채광", "level": 0, "location": "광산", "is_mining": True, "type": "채광"}
    item_to_boss["은광석"] = {"boss": "광산-채광", "level": 0, "location": "광산", "is_mining": True, "type": "채광"}
    item_to_boss["금광석"] = {"boss": "광산-채광", "level": 0, "location": "광산", "is_mining": True, "type": "채광"}
    item_to_boss["맑은 샘의 보석"] = {"boss": "멀록해변-채광", "level": 0, "location": "멀록해변", "is_mining": True, "type": "채광"}
    item_to_boss["현빙"] = {"boss": "운설산-채광", "level": 0, "location": "운설산", "is_mining": True, "type": "채광"}
    item_to_boss["지옥석"] = {"boss": "용암화산-채광", "level": 0, "location": "용암화산", "is_mining": True, "type": "채광"}

    sorted_bosses = sorted(list(boss_map.values()), key=lambda x: (x["level"], x["name"]))
    boss_loc_lookup = {b["name"]: b["location"] for b in sorted_bosses}

    # 2. Parse 6 Category Sheets: 무기, 갑옷, 투구, 장신구, 보조장비, 마나석
    categories = ["무기", "갑옷", "투구", "장신구", "보조장비", "마나석"]
    category_gear = {cat: {"top_gear": [], "recipes": []} for cat in categories}
    global_recipe_map = {}

    # Pass 1: Parse recipes from category sheets
    raw_cat_recipes = {}
    for cname in categories:
        if cname not in wb.sheetnames:
            continue
        s = wb[cname]
        rows = list(s.iter_rows(values_only=True))
        curr_item = ""
        curr_rc = None
        curr_sub_cat = ""
        c_recipes = []
        
        for r in rows[1:]:
            if not any(r):
                continue
            sub_cat = clean_str(r[0])
            if sub_cat:
                curr_sub_cat = sub_cat.replace(" ", "")
            it_name = clean_str(r[1])
            mat_name = clean_str(r[2])
            qty = r[3] if len(r) > 3 and r[3] is not None else 1
            lvl = clean_lvl(r[4]) if len(r) > 4 else ""
            synergy = clean_str(r[5]) if len(r) > 5 else ""
            
            if it_name and it_name != curr_item:
                curr_item = it_name
                synergy_jobs = []
                if synergy:
                    parts = re.split(r'[,/ ]+', synergy)
                    synergy_jobs = [clean_str(p) for p in parts if clean_str(p) in ALL_31_JOBS]
                
                curr_rc = {
                    "name": curr_item,
                    "category": cname,
                    "sub_cat": curr_sub_cat if curr_sub_cat else cname,
                    "level": int(lvl) if lvl.isdigit() else 0,
                    "level_str": f"Lv.{lvl}" if lvl else "",
                    "synergy": synergy,
                    "synergy_jobs": synergy_jobs,
                    "special_effect": synergy,
                    "materials": [],
                    "is_drop": False,
                    "drop_boss": "",
                    "drop_level": 0,
                    "drop_location": "",
                    "drop_type": "아이템"
                }
                c_recipes.append(curr_rc)
                global_recipe_map[curr_item] = curr_rc
                
            if mat_name and curr_rc:
                q_val = 1
                if isinstance(qty, int):
                    q_val = qty
                elif str(qty).isdigit():
                    q_val = int(qty)
                elif "*" in mat_name:
                    parts = mat_name.split("*")
                    mat_name = clean_str(parts[0])
                    if len(parts) > 1 and parts[1].strip().isdigit():
                        q_val = int(parts[1].strip())
                        
                curr_rc["materials"].append({
                    "name": mat_name,
                    "qty": q_val,
                    "boss": "",
                    "level": 0,
                    "location": "",
                    "level_str": ""
                })

        # Mark drop items (items with 0 materials that drop from bosses)
        for rc in c_recipes:
            if len(rc["materials"]) == 0:
                it_lookup = clean_str(rc["name"])
                if it_lookup in item_to_boss:
                    drop_info = item_to_boss[it_lookup]
                    rc["is_drop"] = True
                    rc["drop_boss"] = drop_info["boss"]
                    rc["drop_level"] = drop_info["level"]
                    rc["drop_location"] = drop_info.get("location", "")
                    rc["drop_type"] = drop_info.get("type", "아이템")
                    if not rc["level"] and drop_info["level"]:
                        rc["level"] = drop_info["level"]
                        rc["level_str"] = f"Lv.{drop_info['level']}"

        raw_cat_recipes[cname] = c_recipes

    # Pass 2: Auto-resolve boss, level, location for materials
    for cname, r_list in raw_cat_recipes.items():
        for rc in r_list:
            for m in rc["materials"]:
                m_name = m["name"]
                c_mat = clean_str(m_name)
                target_rc = global_recipe_map.get(m_name) or global_recipe_map.get(c_mat)

                if target_rc and len(target_rc.get("materials", [])) > 0 and not target_rc.get("is_drop"):
                    # Regular craftable combo item
                    m["boss"] = "조합템"
                    m["level"] = target_rc.get("level", 0)
                    m["level_str"] = f"Lv.{m['level']}" if m["level"] else ""
                elif target_rc and target_rc.get("is_drop"):
                    # Intermediate item that is directly dropped from a boss
                    m["boss"] = target_rc["drop_boss"]
                    m["level"] = target_rc["drop_level"] or target_rc["level"]
                    m["location"] = target_rc["drop_location"]
                    m["level_str"] = f"Lv.{m['level']}" if m["level"] > 0 else ""
                    m["is_drop"] = True
                elif m_name in item_to_boss or c_mat in item_to_boss:
                    # Material in boss drop table
                    info = item_to_boss.get(m_name) or item_to_boss.get(c_mat)
                    m["boss"] = info["boss"]
                    m["level"] = info["level"]
                    m["location"] = info.get("location", "")
                    m["level_str"] = f"Lv.{info['level']}" if info["level"] > 0 else ("채광" if "채광" in m["boss"] else "")
                    m["is_drop"] = True
                else:
                    m["boss"] = "조합템"

    # Pass 3: Build top_gear per category
    # Include end-game crafted items (not in mats_used) first sorted by level desc,
    # followed by all other gear (including drop gear like 천둥의 격노) sorted by level desc.
    for cname, r_list in raw_cat_recipes.items():
        mats_used = set()
        for rc in r_list:
            for m in rc["materials"]:
                mats_used.add(m["name"])
                mats_used.add(clean_str(m["name"]))

        end_game_items = [rc for rc in r_list if rc["name"] not in mats_used]
        end_game_items.sort(key=lambda x: x["level"], reverse=True)

        other_items = [rc for rc in r_list if rc["name"] in mats_used]
        other_items.sort(key=lambda x: x["level"], reverse=True)

        # All items are selectable, with end-game items prioritised
        all_ordered = [rc["name"] for rc in end_game_items] + [rc["name"] for rc in other_items]

        category_gear[cname]["recipes"] = r_list
        category_gear[cname]["top_gear"] = all_ordered

    # 3. Parse Job Sheets (드루이드, 프리스트, 가디언, 블레이드 스피릿, 얼음마법사, 검사)
    job_sheets = ["프리스트", "가디언", "드루이드", "블레이드 스피릿", "얼음마법사", "검사"]
    legacy_jobs = {}
    
    for jname in job_sheets:
        if jname not in wb.sheetnames:
            continue
        sheet = wb[jname]
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < 2:
            continue
            
        h_idx = 0
        for idx, r in enumerate(rows):
            r_str = [str(cell) for cell in r if cell is not None]
            if any("아이템" in x or "재료" in x or "최종템" in x for x in r_str):
                h_idx = idx
                break
                
        col_map = {}
        for idx, col in enumerate(rows[h_idx]):
            c = clean_str(col)
            if c in ["종류", "분류", "부위"]:
                col_map["cat"] = idx
            elif c in ["아이템 이름", "아이템", "최종템", "장비"]:
                col_map["item"] = idx
            elif c in ["재료", "필요재료"]:
                col_map["mat"] = idx
            elif c in ["개수", "수량"]:
                col_map["qty"] = idx
            elif c in ["보스", "드랍보스", "획득처"]:
                col_map["boss"] = idx
            elif c in ["레벨", "Lv", "렙"]:
                col_map["lvl"] = idx
                
        c_cat = col_map.get("cat", 0)
        c_item = col_map.get("item", 1)
        c_mat = col_map.get("mat", 2)
        c_qty = col_map.get("qty", 3)
        c_boss = col_map.get("boss", 4)
        c_lvl = col_map.get("lvl", 5)
        
        j_recipes = []
        curr_cat = ""
        curr_item = ""
        curr_rc = None
        
        for r in rows[h_idx + 1:]:
            cat = clean_str(r[c_cat]) if len(r) > c_cat else ""
            item = clean_str(r[c_item]) if len(r) > c_item else ""
            mat = clean_str(r[c_mat]) if len(r) > c_mat else ""
            qty = clean_str(r[c_qty]) if len(r) > c_qty else "1"
            boss = clean_str(r[c_boss]) if len(r) > c_boss else ""
            lvl = clean_lvl(r[c_lvl]) if len(r) > c_lvl else ""
            
            if cat:
                curr_cat = cat
            if item and item != curr_item:
                curr_item = item
                base_rc = global_recipe_map.get(curr_item) or global_recipe_map.get(clean_str(curr_item))
                item_lvl = int(lvl) if (lvl and lvl.isdigit()) else (base_rc.get("level", 0) if base_rc else 0)
                if not item_lvl and curr_item == "신앙의 갑옷":
                    item_lvl = 380
                elif not item_lvl and curr_item in item_to_boss:
                    item_lvl = item_to_boss[curr_item]["level"]
                elif not item_lvl and clean_str(curr_item) in item_to_boss:
                    item_lvl = item_to_boss[clean_str(curr_item)]["level"]
                item_lvl_str = f"Lv.{item_lvl}" if item_lvl > 0 else ""
                item_synergy = base_rc.get("synergy", "") if base_rc else ""
                item_synergy_jobs = base_rc.get("synergy_jobs", []) if base_rc else []

                curr_rc = {
                    "category": curr_cat or (base_rc.get("category", "") if base_rc else ""),
                    "sub_cat": (base_rc.get("sub_cat", "") if base_rc else "") or curr_cat,
                    "name": curr_item,
                    "level": item_lvl,
                    "level_str": item_lvl_str,
                    "synergy": item_synergy,
                    "synergy_jobs": item_synergy_jobs,
                    "special_effect": item_synergy,
                    "materials": []
                }
                j_recipes.append(curr_rc)
                if curr_item not in global_recipe_map:
                    global_recipe_map[curr_item] = curr_rc
                    
            if mat and curr_rc:
                q_val = 1
                if qty and qty.isdigit():
                    q_val = int(qty)
                elif "*" in mat:
                    parts = mat.split("*")
                    mat = clean_str(parts[0])
                    if len(parts) > 1 and parts[1].strip().isdigit():
                        q_val = int(parts[1].strip())
                        
                if not boss or boss in ["재료", "아이템"]:
                    if mat in global_recipe_map:
                        boss = "조합템"
                    elif mat in item_to_boss:
                        boss = item_to_boss[mat]["boss"]
                        if not lvl and item_to_boss[mat]["level"] > 0:
                            lvl = str(item_to_boss[mat]["level"])
                    else:
                        boss = "조합템"
                        
                loc = boss_loc_lookup.get(boss, "")
                lvl_int = int(lvl) if lvl.isdigit() else 0
                lvl_str = f"Lv.{lvl}" if (lvl.isdigit() and int(lvl) > 0) else ("채광" if "채광" in boss else "")
                
                curr_rc["materials"].append({
                    "name": mat,
                    "qty": q_val,
                    "boss": boss if boss else "조합템",
                    "level": lvl_int,
                    "level_str": lvl_str,
                    "location": loc
                })
                
        used_as_mat = set()
        for r in j_recipes:
            for m in r["materials"]:
                if m["boss"] == "조합템" or m["name"] in global_recipe_map:
                    used_as_mat.add(m["name"])
                    
        gear_slots = {cat: [] for cat in categories}
        for r in j_recipes:
            if r["name"] not in used_as_mat:
                c = r["category"]
                if c in gear_slots:
                    gear_slots[c].append(r["name"])
                    
        legacy_jobs[jname] = {
            "name": jname,
            "recipes": j_recipes,
            "recipe_map": {r["name"]: r for r in j_recipes},
            "gear_slots": gear_slots
        }

    # 4. Build job synergy recommendations and default loadouts for ALL 31 jobs
    jobs = {}
    for job_name in ALL_31_JOBS:
        recomms = {cat: [] for cat in categories}
        for cat in categories:
            for rc in category_gear[cat]["recipes"]:
                if job_name in rc.get("synergy_jobs", []):
                    recomms[cat].append({
                        "name": rc["name"],
                        "level": rc["level"],
                        "synergy": rc["synergy"],
                        "sub_cat": rc["sub_cat"]
                    })
                    
        default_loadout = {}
        if job_name in legacy_jobs:
            for cat in categories:
                top_items = legacy_jobs[job_name]["gear_slots"].get(cat, [])
                if top_items:
                    default_loadout[cat] = top_items[0]
                elif category_gear[cat]["top_gear"]:
                    default_loadout[cat] = category_gear[cat]["top_gear"][0]
        else:
            for cat in categories:
                cat_top = category_gear[cat]["top_gear"]
                job_recs = [r["name"] for r in recomms[cat] if r["name"] in cat_top]
                if job_recs:
                    default_loadout[cat] = job_recs[0]
                elif cat_top:
                    default_loadout[cat] = cat_top[0]
                else:
                    default_loadout[cat] = ""
                    
        gear_slots = {cat: [default_loadout[cat]] if default_loadout.get(cat) else [] for cat in categories}
        
        active_recipes = []
        visited = set()
        def collect_recipes(item_name):
            if item_name in visited:
                return
            visited.add(item_name)
            rc = global_recipe_map.get(item_name)
            if rc:
                active_recipes.append(rc)
                for m in rc["materials"]:
                    if m["name"] in global_recipe_map:
                        collect_recipes(m["name"])
                        
        for slot_item in default_loadout.values():
            if slot_item:
                collect_recipes(slot_item)
                
        active_recipe_map = {r["name"]: r for r in active_recipes}
        
        jobs[job_name] = {
            "name": job_name,
            "recipes": active_recipes,
            "recipe_map": active_recipe_map,
            "gear_slots": gear_slots,
            "default_loadout": default_loadout,
            "recommendations": recomms
        }
        
    return {
        "all_job_list": ALL_31_JOBS,
        "job_list": ALL_31_JOBS,
        "bosses": sorted_bosses,
        "jobs": jobs,
        "category_gear": category_gear,
        "global_recipe_map": global_recipe_map,
        "legacy_jobs": legacy_jobs
    }

if __name__ == "__main__":
    os.makedirs(os.path.join(BASE_DIR, 'data'), exist_ok=True)
    data = parse_excel(EXCEL_PATH)
    
    with open(OUTPUT_JSON, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    with open(OUTPUT_JS, 'w', encoding='utf-8') as f:
        f.write("window.PRELOADED_GAME_DATA = " + json.dumps(data, ensure_ascii=False, indent=2) + ";\n")

    print(f"Successfully generated {OUTPUT_JSON} and {OUTPUT_JS}")
    print(f"Total jobs: {len(data['jobs'])}")
    print(f"Total bosses: {len(data['bosses'])}")
    print(f"Total global recipes: {len(data['global_recipe_map'])}")
    for cat in ["무기", "갑옷", "투구", "장신구", "보조장비", "마나석"]:
        print(f"  [{cat}]: {len(data['category_gear'][cat]['top_gear'])} top gear, {len(data['category_gear'][cat]['recipes'])} total recipes")
